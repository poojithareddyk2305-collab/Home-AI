"""Turns the model's Markdown into a formatted Word (.docx) file."""
import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

INLINE = re.compile(r"(\*\*[^*]+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`|\[[^\]]+?\]\(https?://[^)\s]+\))")
LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
HRULE = re.compile(r"^\s*([-*_])\1{2,}\s*$")


def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _hyperlink(paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E9E")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str, bold: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        link = LINK.fullmatch(part)
        if link:
            _hyperlink(paragraph, link.group(1), link.group(2))
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.bold = bold or None
        else:
            run = paragraph.add_run(part)
            run.bold = bold or None


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _add_table(doc, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, row[j] if j < len(row) else "", bold=(i == 0))
            for run in paragraph.runs:
                run.font.size = Pt(10)
            if i == 0:
                _shade(cell._tc.get_or_add_tcPr(), "E8ECF7")
    doc.add_paragraph()


def markdown_to_docx(title: str, markdown: str) -> bytes:
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    doc.core_properties.title = title
    doc.core_properties.author = "AI Chatbot"

    lines = markdown.replace("\r\n", "\n").split("\n")
    # the title is rendered separately, so drop a duplicate leading "# Title"
    if lines and HEADING.match(lines[0]) and HEADING.match(lines[0]).group(1) == "#":
        if HEADING.match(lines[0]).group(2).strip().lower() == title.strip().lower():
            lines = lines[1:]

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    i = 0
    para_buffer: list[str] = []

    def flush_paragraph():
        if para_buffer:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para_buffer))
            para_buffer.clear()

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            flush_paragraph()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            for code_line in code or [""]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(code_line)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                _shade(p._p.get_or_add_pPr(), "F1F1F6")
            i += 1
            continue

        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        heading_match = HEADING.match(line)
        if heading_match:
            flush_paragraph()
            doc.add_heading(heading_match.group(2), level=min(len(heading_match.group(1)), 4))
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            flush_paragraph()
            rows = [_split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue

        if HRULE.match(line):
            flush_paragraph()
            i += 1
            continue

        bullet = BULLET.match(line)
        numbered = NUMBERED.match(line)
        if bullet or numbered:
            flush_paragraph()
            m = bullet or numbered
            depth = min(len(m.group(1).replace("\t", "  ")) // 2, 2)
            base_style = "List Bullet" if bullet else "List Number"
            style = base_style if depth == 0 else f"{base_style} {depth + 1}"
            add_inline(doc.add_paragraph(style=style), m.group(2))
            i += 1
            continue

        if line.lstrip().startswith(">"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, line.lstrip()[1:].strip())
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x66)
            i += 1
            continue

        para_buffer.append(line)
        i += 1

    flush_paragraph()

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    run = footer.add_run(f"Generated by AI Chatbot on {datetime.now().strftime('%d %B %Y')}")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x99)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
